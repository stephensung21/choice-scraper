"""Scrape CHOICE (choice.com.au) product reviews for a few categories.

For personal use with your own CHOICE membership. You log in yourself in a
real browser window; the session is kept in .browser-profile/ so later runs
reuse it. Requests are spaced out to stay gentle on the site.

Usage:
    python scrape.py login                    # log in once, by hand
    python scrape.py scrape                   # scrape every category in categories.json
    python scrape.py scrape kettles tvs       # scrape only some categories
    python scrape.py scrape --limit 3         # first 3 products per category (good for a test run)
    python scrape.py export                   # rebuild the Excel file from saved data
    python scrape.py debug <url>              # save a page's HTML + screenshot for troubleshooting
"""

import argparse
import hashlib
import json
import os
import random
import re
import sys
import time
from pathlib import Path
from urllib.parse import urljoin, urlparse

from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent
PROFILE_DIR = ROOT / ".browser-profile"
CATEGORIES_FILE = ROOT / "categories.json"
OUTPUT_DIR = ROOT / "output"
PRODUCTS_DIR = OUTPUT_DIR / "products"
DEBUG_DIR = ROOT / "debug"
HOME_URL = "https://www.choice.com.au/"

MAX_LISTING_PAGES = 30
MAX_LOAD_MORE_CLICKS = 40
EXCEL_CELL_LIMIT = 32000

# Runs inside the page: collects every product link on a listing page that
# points deeper than the listing URL itself (or matches link_pattern if given).
COLLECT_LINKS_JS = """
([listingPath, linkPattern]) => {
  const re = linkPattern ? new RegExp(linkPattern) : null;
  const base = listingPath.replace(/\\/$/, '');
  const links = new Set();
  for (const a of document.querySelectorAll('a[href]')) {
    let u;
    try { u = new URL(a.getAttribute('href'), location.href); } catch (e) { continue; }
    if (u.hostname !== location.hostname) continue;
    const path = u.pathname.replace(/\\/$/, '');
    const ok = re ? re.test(u.href) : (path.startsWith(base + '/') && path.length > base.length + 1);
    if (ok) links.add(u.origin + path);
  }
  return [...links];
}
"""

# Runs inside the page: pulls out everything that looks like product detail.
EXTRACT_PRODUCT_JS = """
() => {
  const clean = s => (s || '').replace(/\\s+/g, ' ').trim();
  const cleanLines = s => (s || '').split('\\n').map(clean).filter(Boolean).join('; ');
  const root = document.querySelector('main') || document.body;
  const out = { url: location.href, title: clean(document.querySelector('h1')?.innerText) };

  out.json_ld = [...document.querySelectorAll('script[type="application/ld+json"]')]
    .map(s => { try { return JSON.parse(s.textContent); } catch (e) { return null; } })
    .filter(Boolean);

  out.meta = {};
  document.querySelectorAll('meta[property], meta[name]').forEach(m => {
    const k = m.getAttribute('property') || m.getAttribute('name');
    const v = m.getAttribute('content');
    if (k && v) out.meta[k] = v;
  });

  const specs = {};
  const add = (k, v) => {
    k = clean(k).replace(/:$/, ''); v = clean(v);
    if (!k || !v || k.length > 120 || v.length > 2000 || k === v) return;
    let key = k, i = 2;
    while (key in specs && specs[key] !== v) key = `${k} (${i++})`;
    specs[key] = v;
  };
  root.querySelectorAll('tr').forEach(tr => {
    const cells = [...tr.children].filter(c => c.tagName === 'TD' || c.tagName === 'TH');
    if (cells.length === 2) add(cells[0].innerText, cells[1].innerText);
  });
  root.querySelectorAll('dl').forEach(dl => {
    let k = null;
    for (const c of dl.children) {
      if (c.tagName === 'DT') k = c.innerText;
      else if (c.tagName === 'DD' && k) add(k, c.innerText);
    }
  });
  root.querySelectorAll('li, p').forEach(el => {
    if (el.children.length > 3) return;
    const m = clean(el.innerText).match(/^([^:]{2,60}):\\s+(.{1,300})$/);
    if (m) add(m[1], m[2]);
  });
  out.specs = specs;

  const sections = {};
  for (const h of root.querySelectorAll('h2, h3, h4')) {
    const name = clean(h.innerText);
    if (!name) continue;
    const parts = [];
    let n = h.nextElementSibling;
    while (n && !/^H[1-4]$/.test(n.tagName) && !n.querySelector('h2, h3, h4')) {
      const t = cleanLines(n.innerText);
      if (t) parts.push(t);
      n = n.nextElementSibling;
    }
    if (parts.length) sections[name] = (sections[name] ? sections[name] + '\\n' : '') + parts.join('\\n');
  }
  out.sections = sections;
  out.full_text = root.innerText;
  return out;
}
"""


def load_categories():
    data = json.loads(CATEGORIES_FILE.read_text())
    return {k: v for k, v in data.items() if not k.startswith("_")}


def open_browser(p, headless=False):
    kwargs = {"headless": headless, "viewport": {"width": 1400, "height": 1000}}
    if os.environ.get("CHOICE_BROWSER_PATH"):
        kwargs["executable_path"] = os.environ["CHOICE_BROWSER_PATH"]
    ctx = p.chromium.launch_persistent_context(str(PROFILE_DIR), **kwargs)
    page = ctx.pages[0] if ctx.pages else ctx.new_page()
    return ctx, page


def polite_pause(delay):
    time.sleep(delay + random.uniform(0, delay))


def goto(page, url):
    page.goto(url, wait_until="domcontentloaded", timeout=60000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except PlaywrightTimeout:
        pass


def save_debug(page, name):
    DEBUG_DIR.mkdir(exist_ok=True)
    (DEBUG_DIR / f"{name}.html").write_text(page.content(), encoding="utf-8")
    page.screenshot(path=str(DEBUG_DIR / f"{name}.png"), full_page=True)
    print(f"    saved debug/{name}.html and debug/{name}.png")


def scroll_and_expand(page):
    """Scroll to the bottom and keep clicking 'Load more'-style buttons."""
    button = page.get_by_role("button", name=re.compile(r"(load|show|view) more", re.I))
    for _ in range(MAX_LOAD_MORE_CLICKS):
        page.mouse.wheel(0, 20000)
        page.wait_for_timeout(800)
        if not button.count() or not button.first.is_visible():
            break
        try:
            button.first.click(timeout=5000)
            page.wait_for_timeout(1500)
        except PlaywrightTimeout:
            break


def next_page_url(page):
    nxt = page.locator('a[rel="next"]')
    if not nxt.count():
        nxt = page.get_by_role("link", name=re.compile(r"^\s*next\b", re.I))
    if nxt.count():
        href = nxt.first.get_attribute("href")
        if href:
            return urljoin(page.url, href)
    return None


def collect_product_links(page, cat_key, cat, delay, interactive):
    listing_path = urlparse(cat["url"]).path
    link_pattern = cat.get("link_pattern")
    links, url, seen_pages = [], cat["url"], set()

    while url and url not in seen_pages and len(seen_pages) < MAX_LISTING_PAGES:
        seen_pages.add(url)
        goto(page, url)
        scroll_and_expand(page)
        for link in page.evaluate(COLLECT_LINKS_JS, [listing_path, link_pattern]):
            if link not in links:
                links.append(link)
        url = next_page_url(page)
        if url:
            polite_pause(delay)

    if not links:
        print(f"    no product links found on {cat['url']}")
        save_debug(page, f"listing-{cat_key}")
        if interactive:
            input(
                "    In the browser window, go to the page that lists the products "
                "for this category, then press Enter here..."
            )
            scroll_and_expand(page)
            listing_path = urlparse(page.url).path
            links = page.evaluate(COLLECT_LINKS_JS, [listing_path, link_pattern])
            if links:
                print(f"    tip: put this URL in categories.json for next time: {page.url}")
    return links


def first(*values):
    for v in values:
        if v not in (None, "", [], {}):
            return v
    return None


def find_jsonld_product(blobs):
    stack = list(blobs)
    while stack:
        item = stack.pop()
        if isinstance(item, list):
            stack.extend(item)
        elif isinstance(item, dict):
            types = item.get("@type")
            types = types if isinstance(types, list) else [types]
            if "Product" in types:
                return item
            stack.extend(item.get("@graph", []))
    return {}


def spec_lookup(specs, pattern):
    for k, v in specs.items():
        if re.search(pattern, k, re.I):
            return v
    return None


def section_lookup(sections, pattern):
    for k, v in sections.items():
        if re.search(pattern, k, re.I):
            return v
    return None


def summarise(raw, cat_key):
    """Derive the headline fields from the raw page extraction."""
    text = raw.get("full_text", "")
    specs, sections = raw.get("specs", {}), raw.get("sections", {})
    product = find_jsonld_product(raw.get("json_ld", []))
    offers = product.get("offers") or {}
    if isinstance(offers, list):
        offers = offers[0] if offers else {}
    brand = product.get("brand")
    if isinstance(brand, dict):
        brand = brand.get("name")

    score_match = re.search(
        r"(?:CHOICE\s+Expert\s+Rating|CHOICE\s+score|Overall\s+score|Expert\s+rating)\D{0,30}?(\d{1,3})\s*%?",
        text, re.I,
    )
    price_match = re.search(r"\$\s?\d[\d,]*(?:\.\d{2})?", text)

    return {
        "category": cat_key,
        "name": first(raw.get("title"), product.get("name"), raw.get("meta", {}).get("og:title")),
        "brand": first(brand, spec_lookup(specs, r"^brand$")),
        "price": first(offers.get("price"), spec_lookup(specs, r"price"),
                       price_match.group(0) if price_match else None),
        "choice_score": first(score_match.group(1) if score_match else None,
                              spec_lookup(specs, r"score|rating")),
        "recommended": bool(re.search(r"\bRecommended\b", text)),
        "best_buy": bool(re.search(r"\bBest\s+Buy\b", text, re.I)),
        "pros": section_lookup(sections, r"^(pros|good|what we like|the good)"),
        "cons": section_lookup(sections, r"^(cons|bad|what we don.t like|the bad)"),
        "verdict": section_lookup(sections, r"verdict|summary|bottom line|our say"),
        "possibly_locked": bool(re.search(r"become a member|join (now|choice)|unlock", text, re.I))
                           and not score_match,
        "url": raw.get("url"),
    }


def product_file(cat_key, url):
    slug = re.sub(r"[^a-z0-9]+", "-", urlparse(url).path.lower()).strip("-")[-80:]
    digest = hashlib.sha1(url.encode()).hexdigest()[:8]
    return PRODUCTS_DIR / cat_key / f"{slug}-{digest}.json"


def scrape_product(page, cat_key, url):
    goto(page, url)
    # Open any collapsed "show more / specifications" panels so their text is captured.
    for el in page.locator("button[aria-expanded='false'], details:not([open]) > summary").all()[:30]:
        try:
            if el.is_visible():
                el.click(timeout=2000)
        except Exception:
            pass
    page.wait_for_timeout(500)
    raw = page.evaluate(EXTRACT_PRODUCT_JS)
    return {"summary": summarise(raw, cat_key), "raw": raw}


def cmd_login(args):
    with sync_playwright() as p:
        ctx, page = open_browser(p)
        goto(page, HOME_URL)
        print("A browser window is open. Log in to CHOICE there (click 'Log in' at the top).")
        input("When you can see you're logged in, press Enter here to save the session... ")
        ctx.close()
    print("Session saved in .browser-profile/ - you can now run: python scrape.py scrape")


def cmd_scrape(args):
    categories = load_categories()
    selected = args.categories or list(categories)
    unknown = [c for c in selected if c not in categories]
    if unknown:
        sys.exit(f"Unknown categories: {unknown}. Options: {list(categories)}")

    with sync_playwright() as p:
        ctx, page = open_browser(p, headless=args.headless)
        for cat_key in selected:
            cat = categories[cat_key]
            print(f"\n== {cat['name']} ==")
            links = collect_product_links(page, cat_key, cat, args.delay, not args.headless)
            print(f"    found {len(links)} product links")
            if args.limit:
                links = links[: args.limit]

            for i, url in enumerate(links, 1):
                path = product_file(cat_key, url)
                if path.exists() and not args.refresh:
                    print(f"    [{i}/{len(links)}] already saved, skipping {url}")
                    continue
                polite_pause(args.delay)
                try:
                    data = scrape_product(page, cat_key, url)
                except Exception as e:
                    print(f"    [{i}/{len(links)}] FAILED {url}: {e}")
                    continue
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
                s = data["summary"]
                flag = "  (looks locked - are you logged in?)" if s["possibly_locked"] else ""
                print(f"    [{i}/{len(links)}] {s['name']} | score {s['choice_score']} | {s['price']}{flag}")
        ctx.close()
    cmd_export(args)


def load_saved():
    rows = {}
    for path in sorted(PRODUCTS_DIR.glob("*/*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        rows.setdefault(path.parent.name, []).append(data)
    return rows


def score_key(row):
    try:
        return -float(row["summary"]["choice_score"])
    except (TypeError, ValueError):
        return float("inf")


def cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and len(value) > EXCEL_CELL_LIMIT:
        value = value[:EXCEL_CELL_LIMIT] + " ...[truncated, see JSON]"
    return value


def cmd_export(args):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    saved = load_saved()
    if not saved:
        print("Nothing scraped yet.")
        return

    wb = Workbook()
    overview = wb.active
    overview.title = "All products"
    summary_cols = list(next(iter(saved.values()))[0]["summary"].keys())
    overview.append(summary_cols)

    for cat_key, items in saved.items():
        items.sort(key=score_key)
        (OUTPUT_DIR / f"{cat_key}.json").write_text(
            json.dumps(items, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        spec_cols = sorted({k for it in items for k in it["raw"].get("specs", {})})
        section_cols = sorted({k for it in items for k in it["raw"].get("sections", {})})

        ws = wb.create_sheet(cat_key[:31])
        ws.append(summary_cols + [f"spec: {c}" for c in spec_cols] + [f"section: {c}" for c in section_cols])
        for it in items:
            s, raw = it["summary"], it["raw"]
            row = [s.get(c) for c in summary_cols]
            row += [raw.get("specs", {}).get(c) for c in spec_cols]
            row += [raw.get("sections", {}).get(c) for c in section_cols]
            ws.append([cell(v) for v in row])
            overview.append([cell(s.get(c)) for c in summary_cols])
        ws.freeze_panes = "C2"

    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)

    OUTPUT_DIR.mkdir(exist_ok=True)
    xlsx = OUTPUT_DIR / "choice_products.xlsx"
    wb.save(xlsx)
    total = sum(len(v) for v in saved.values())
    print(f"\nWrote {total} products to {xlsx.relative_to(ROOT)} (plus one JSON file per category in output/)")


def cmd_debug(args):
    with sync_playwright() as p:
        ctx, page = open_browser(p, headless=args.headless)
        goto(page, args.url)
        save_debug(page, "page")
        data = page.evaluate(EXTRACT_PRODUCT_JS)
        (DEBUG_DIR / "page-extract.json").write_text(json.dumps(data, indent=2, ensure_ascii=False))
        print("    saved debug/page-extract.json")
        ctx.close()


def main():
    parser = argparse.ArgumentParser(description="Scrape CHOICE product reviews (personal use).")
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("login", help="open a browser so you can log in once")

    sp = sub.add_parser("scrape", help="scrape product reviews")
    sp.add_argument("categories", nargs="*", help="category keys from categories.json (default: all)")
    sp.add_argument("--limit", type=int, help="max products per category")
    sp.add_argument("--delay", type=float, default=3.0, help="base seconds between page loads (default 3)")
    sp.add_argument("--refresh", action="store_true", help="re-scrape products already saved")
    sp.add_argument("--headless", action="store_true", help="hide the browser window")

    sub.add_parser("export", help="rebuild output/choice_products.xlsx from saved data")

    dp = sub.add_parser("debug", help="save a page's HTML, screenshot and extraction to debug/")
    dp.add_argument("url")
    dp.add_argument("--headless", action="store_true")

    args = parser.parse_args()
    {"login": cmd_login, "scrape": cmd_scrape, "export": cmd_export, "debug": cmd_debug}[args.command](args)


if __name__ == "__main__":
    main()
